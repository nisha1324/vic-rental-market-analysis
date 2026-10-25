"""Turn the Rental Report's wide Excel sheets into tidy (long) tables.

Every time-series sheet in these workbooks has the same shape:

    row 0      title
    row 1      dwelling type, then one quarter label per measure column
    row 2      measure names ("Count"/"Median" or "Affordable"/"Percent")
    row 3..    label column(s), then values; '-' means suppressed (too few
               lettings to publish), blank means no value

The suburb and LGA sheets have two label columns (region, area). The region
is only written on the first row of each block, so it is filled down. The
affordability sheets have one label column (the LGA) and their header rows
sit one row lower.

Anything unexpected in a value cell (a stray note, a new symbol) raises
instead of turning silently into a missing value.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

import openpyxl
import pandas as pd

QUARTER_RE = re.compile(r"^(Mar|Jun|Sep|Dec) (\d{4})$")
QUARTER_END_MONTH = {"Mar": 3, "Jun": 6, "Sep": 9, "Dec": 12}
SUPPRESSED = "-"

# Spelling differences between workbooks (and typos in the source), mapped
# to one name so the tables join. Applied to suburb and LGA labels.
NAME_FIXES = {
    "Wanagaratta": "Wangaratta",
    "Geelong-Newcomb": "Geelong-Newcombe",
    "Mornington Penin'a": "Mornington Peninsula",
    "Queenscliff": "Queenscliffe",
}

# Suburb workbook regions outside metropolitan Melbourne (the report's own
# split: Table 4 lists these towns under "Regional Victoria").
REGIONAL_SUBURB_REGIONS = {"Geelong", "Ballarat", "Bendigo", "Other Regional Centres"}

SUBURB_SHEETS = {
    "1 bedroom flat": "1br_flat",
    "2 bedroom flat": "2br_flat",
    "3 bedroom flat": "3br_flat",
    "2 bedroom house": "2br_house",
    "3 bedroom house": "3br_house",
    "4 bedroom house": "4br_house",
    "All properties": "all",
}
LGA_SHEETS = {
    "1br flat": "1br_flat",
    "2br Flat": "2br_flat",
    "3br Flat": "3br_flat",
    "2br House": "2br_house",
    "3br House": "3br_house",
    "4br House": "4br_house",
    "All Properties": "all",
}
AFFORDABILITY_SHEETS = {
    "lga aff 1br": "1br",
    "lga aff 2br": "2br",
    "lga aff 3br": "3br",
    "lga aff 4br": "4br",
    "lga aff total": "all",
}

# How far a repeated Victoria total row may drift from the first one
# (rounding in the source; see _drop_duplicate and the data audit).
LGA_REPEAT_TOLERANCE = {"lettings": 1, "median_rent": 0}
AFFORDABILITY_REPEAT_TOLERANCE = {"affordable_lettings": 2, "affordable_share": 0.001}


def quarter_end(label: str) -> pd.Timestamp:
    """'Sep 2025' -> Timestamp('2025-09-30')."""
    m = QUARTER_RE.match(label.strip())
    if not m:
        raise ValueError(f"not a quarter label: {label!r}")
    month, year = QUARTER_END_MONTH[m.group(1)], int(m.group(2))
    return pd.Timestamp(year=year, month=month, day=1) + pd.offsets.MonthEnd(0)


def clean_value(value, where: str) -> float:
    if value is None or value == SUPPRESSED:
        return float("nan")
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"unexpected value {value!r} at {where}")
    return float(value)


def clean_name(name: str) -> str:
    name = " ".join(str(name).split())
    return NAME_FIXES.get(name, name)


@dataclass(frozen=True)
class Layout:
    header_row: int      # row holding the quarter labels
    measure_row: int     # row holding the measure names
    label_cols: int      # leading label columns (1 or 2)
    measures: dict       # source measure name -> output column


SUBURB_LAYOUT = Layout(1, 2, 2, {"Count": "lettings", "Median": "median_rent"})
LGA_LAYOUT = Layout(1, 2, 2, {"Count": "lettings", "Median": "median_rent"})
AFFORDABILITY_LAYOUT = Layout(
    2, 3, 1, {"Affordable": "affordable_lettings", "Percent": "affordable_share"}
)


def wide_to_long(rows: list[tuple], layout: Layout, sheet: str) -> pd.DataFrame:
    """One sheet's rows -> long table: labels..., quarter, one column per measure."""
    header, measures = rows[layout.header_row], rows[layout.measure_row]
    columns = []  # (col index, quarter, output measure)
    for j in range(layout.label_cols, len(header)):
        label = header[j]
        if label is None or not str(label).strip():
            continue
        if measures[j] not in layout.measures:
            raise ValueError(f"{sheet}: unknown measure {measures[j]!r} in column {j}")
        columns.append((j, quarter_end(str(label)), layout.measures[measures[j]]))
    if not columns:
        raise ValueError(f"{sheet}: no quarter columns found")

    records = []
    region = None
    for i in range(layout.measure_row + 1, len(rows)):
        row = rows[i]
        labels = [row[k] if k < len(row) else None for k in range(layout.label_cols)]
        if all(v is None for v in labels):
            continue
        if layout.label_cols == 2:
            if labels[0] is not None:
                region = " ".join(str(labels[0]).split())
            area = labels[1]
            key = {"region": region, "area": clean_name(area) if area is not None else None}
        else:
            key = {"area": clean_name(labels[0])}
        values = {}
        for j, quarter, measure in columns:
            cell = row[j] if j < len(row) else None
            values.setdefault(quarter, {})[measure] = clean_value(cell, f"{sheet}!r{i}c{j}")
        for quarter, measured in values.items():
            records.append({**key, "row": i, "quarter": quarter, **measured})
    return pd.DataFrame.from_records(records)


def read_sheets(path: Path, sheets: dict, layout: Layout) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing = set(sheets) - set(wb.sheetnames)
    if missing:
        raise ValueError(f"{path.name}: missing sheets {sorted(missing)}")
    frames = []
    for sheet, dwelling in sheets.items():
        rows = list(wb[sheet].iter_rows(values_only=True))
        df = wide_to_long(rows, layout, sheet)
        df.insert(0, "dwelling", dwelling)
        frames.append(df)
    wb.close()
    return pd.concat(frames, ignore_index=True)


def raw_labels(path: Path) -> set[str]:
    """Every text value in the first two columns of a workbook, as written
    (before NAME_FIXES), so the audit can show which fixes were needed."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    labels = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_col=2, values_only=True):
            labels.update(" ".join(c.split()) for c in row if isinstance(c, str))
    wb.close()
    return labels


def parse_suburbs(path: Path) -> pd.DataFrame:
    """Moving annual median rent and new lettings by suburb group and dwelling type.

    Each value covers leases that started in the 12 months ending that quarter.
    """
    df = read_sheets(path, SUBURB_SHEETS, SUBURB_LAYOUT)
    df["row_type"] = df["area"].map(lambda a: "group_total" if a == "Group Total" else "suburb")
    df["metro"] = ~df["region"].isin(REGIONAL_SUBURB_REGIONS)
    df = df.rename(columns={"area": "suburb"})
    cols = ["dwelling", "region", "suburb", "row_type", "metro", "quarter", "lettings", "median_rent"]
    return _check_unique(df[cols], ["dwelling", "region", "suburb", "quarter"])


def parse_lgas(path: Path) -> pd.DataFrame:
    """Quarterly median rent and new lettings by Local Government Area.

    Below the LGA blocks the sheet has a state total ('Table Total' /
    'Victoria') and a 'METRO NON-METRO' block with Metro, Non-Metro and
    Victoria again. They are kept, tagged by row_type, and the repeated
    Victoria row is dropped after checking it matches (to within one letting).
    """
    df = read_sheets(path, LGA_SHEETS, LGA_LAYOUT)

    def row_type(r) -> str:
        if r.region == "Table Total":
            return "victoria"
        if r.region == "METRO NON-METRO":
            return {"Metro": "metro", "Non-Metro": "non_metro", "Victoria": "victoria_dup"}[r.area]
        return "group_total" if r.area == "Group Total" else "lga"

    df["row_type"] = df.apply(row_type, axis=1)
    df = _drop_duplicate(df, "victoria", "victoria_dup", LGA_REPEAT_TOLERANCE)
    df = df.rename(columns={"area": "lga"})
    cols = ["dwelling", "region", "lga", "row_type", "quarter", "lettings", "median_rent"]
    return _check_unique(df[cols], ["dwelling", "region", "lga", "quarter"])


def parse_affordability(path: Path) -> pd.DataFrame:
    """New lettings affordable to lower-income households, by LGA.

    affordable_share is the share of all new lettings in that LGA and quarter
    that cost no more than 30% of gross income for a household on Centrelink
    payments (Homes Victoria's definition). The sheet ends with 'Table Total'
    (Victoria), Metro, Non-Metro and 'Table Total' again.
    """
    df = read_sheets(path, AFFORDABILITY_SHEETS, AFFORDABILITY_LAYOUT)
    totals = df["area"] == "Table Total"
    first_total = df["dwelling"].map(df[totals].groupby("dwelling")["row"].min())
    df["row_type"] = "lga"
    df.loc[totals, "row_type"] = "victoria_dup"
    df.loc[totals & (df["row"] == first_total), "row_type"] = "victoria"
    df.loc[df["area"] == "Metro", "row_type"] = "metro"
    df.loc[df["area"] == "Non-Metro", "row_type"] = "non_metro"
    df = _drop_duplicate(df, "victoria", "victoria_dup", AFFORDABILITY_REPEAT_TOLERANCE)
    df.loc[df["row_type"] == "victoria", "area"] = "Victoria"
    df = df.rename(columns={"area": "lga"})
    cols = ["dwelling", "lga", "row_type", "quarter", "affordable_lettings", "affordable_share"]
    return _check_unique(df[cols], ["dwelling", "lga", "quarter"])


def parse_table4(path: Path) -> pd.DataFrame:
    """Table 4 of the report: the highest and lowest moving annual median rents
    for 2-bedroom flats and 3-bedroom houses, metro and regional, in the latest
    quarter. Used only to reconcile the parsed suburb data against what Homes
    Victoria published."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Table 4"].iter_rows(values_only=True))
    wb.close()
    dwellings = {"2 bedroom flat": "2br_flat", "3 bedroom house": "3br_house"}
    groups = {"Metropolitan Melbourne": True, "Regional Victoria": False}
    records, dwelling, metro = [], None, None
    for row in rows:
        row = list(row) + [None] * (5 - len(row))
        if row[0] in dwellings:
            dwelling = dwellings[row[0]]
        elif row[0] in groups or row[3] in groups:
            metro = groups[row[0] if row[0] in groups else row[3]]
        else:
            for side, (name, value) in {"highest": (row[0], row[1]), "lowest": (row[3], row[4])}.items():
                if isinstance(value, (int, float)) and name:
                    records.append({"dwelling": dwelling, "metro": metro, "side": side,
                                    "suburb": clean_name(name), "published_median": float(value)})
    return pd.DataFrame.from_records(records)


TABLE14_BEDROOMS = {"1 Bedroom": "1br", "2 Bedrooms": "2br", "3 Bedrooms": "3br",
                    "4+ Bedrooms": "4br", "Total": "all"}
TABLE9_BEDROOMS = {"1 bedroom": "1br", "2 bedroom": "2br", "3 bedroom": "3br", "4+ bedroom": "4br"}
TABLE9_AREAS = {"Metropolitan Melbourne": "metro", "Regional Victoria": "non_metro", "Victoria": "victoria"}
FIG8_AREAS = {"Victoria": "victoria", "Metro": "metro", "Regional": "non_metro"}


def _sheet_rows(path: Path, sheet: str) -> list[list]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"{path.name}: missing sheet {sheet!r}")
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    return rows


def _label(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def parse_fig8(path: Path) -> pd.DataFrame:
    """Source data for Figure 8 of the report: the share of new lettings that
    are affordable, Victoria / Metro / Regional, for the last 21 quarters.

    It overlaps the affordability workbook, so it shows whether past quarters
    were revised. Its last date is also the report's quarter, which the
    Table 9 and Table 14 snapshots don't state themselves."""
    rows = _sheet_rows(path, "Fig 8 source")
    header = next(i for i, r in enumerate(rows) if _label(r[1]).startswith("Victoria"))
    cols = {j: FIG8_AREAS[_label(c).split(" ")[0]] for j, c in enumerate(rows[header])
            if _label(c).split(" ")[0] in FIG8_AREAS}
    if sorted(cols.values()) != sorted(FIG8_AREAS.values()):
        raise ValueError(f"Fig 8 source: unexpected header {rows[header]}")
    records = []
    for i, r in enumerate(rows[header + 1:], start=header + 1):
        if r[0] is None:
            continue
        if not hasattr(r[0], "year"):
            raise ValueError(f"Fig 8 source: unexpected date {r[0]!r} in row {i}")
        quarter = pd.Timestamp(r[0]) + pd.offsets.QuarterEnd(0)
        for j, area in cols.items():
            records.append({"area": area, "quarter": quarter,
                            "affordable_share": clean_value(r[j], f"Fig 8 source!r{i}c{j}")})
    return _check_unique(pd.DataFrame.from_records(records), ["area", "quarter"])


def report_quarter(path: Path) -> pd.Timestamp:
    return parse_fig8(path)["quarter"].max()


def parse_table9(path: Path) -> pd.DataFrame:
    """Table 9: the indicative household for each bedroom count, its weekly
    income, the most it can pay (30% of income) and how many new lettings in
    the report's quarter were at or below that rent, by area."""
    rows = _sheet_rows(path, "Table 9")
    at = lambda name: next(i for i, r in enumerate(rows) if _label(r[0]) == name)  # noqa: E731
    sizes, households = rows[at("Assumed property size")], rows[at("Household type")]
    income, rent = rows[at("Weekly income (net of RA)")], rows[at("Affordable weekly rent")]
    blocks = {"affordable_lettings": at("Affordable rentals (number)"),
              "affordable_share": at("Affordable rentals (% of total)")}
    # each block: a heading row, then one row per area with its name in column 1
    area_rows = {m: {TABLE9_AREAS[_label(rows[i][1])]: i for i in range(start + 1, start + 1 + len(TABLE9_AREAS))
                     if _label(rows[i][1]) in TABLE9_AREAS} for m, start in blocks.items()}
    for m, found in area_rows.items():
        if sorted(found) != sorted(TABLE9_AREAS.values()):
            raise ValueError(f"Table 9: {m} block has areas {sorted(found)}")
    columns = {j: TABLE9_BEDROOMS[_label(s)] for j, s in enumerate(sizes) if _label(s) in TABLE9_BEDROOMS}
    columns[next(j for j, h in enumerate(households) if _label(h) == "Total")] = "all"
    if sorted(columns.values()) != sorted([*TABLE9_BEDROOMS.values(), "all"]):
        raise ValueError(f"Table 9: unexpected property sizes {sizes}")
    quarter = report_quarter(path)
    records = []
    for j, dwelling in columns.items():
        total = dwelling == "all"
        for area in TABLE9_AREAS.values():
            records.append({
                "dwelling": dwelling, "household": "all" if total else _label(households[j]),
                "weekly_income": float("nan") if total else clean_value(income[j], f"Table 9!income c{j}"),
                "affordable_rent": float("nan") if total else clean_value(rent[j], f"Table 9!rent c{j}"),
                "area": area, "quarter": quarter,
                **{m: clean_value(rows[area_rows[m][area]][j], f"Table 9!r{area_rows[m][area]}c{j}")
                   for m in blocks}})
    return _check_unique(pd.DataFrame.from_records(records), ["dwelling", "area"])


def parse_table14(path: Path) -> pd.DataFrame:
    """Table 14: affordable new lettings by LGA and bedroom count in the
    report's quarter (count and share), with the report's Metro/Regional flag.
    The same measures as the affordability workbook, one quarter later."""
    rows = _sheet_rows(path, "Table 14")
    head = next(i for i, r in enumerate(rows) if _label(r[1]) == "1 Bedroom")
    measure_row = next(i for i in range(head, len(rows)) if _label(rows[i][1]) == "#")
    cols = {}  # dwelling -> (count column, share column)
    for j, c in enumerate(rows[head]):
        if _label(c) in TABLE14_BEDROOMS:
            if (_label(rows[measure_row][j]), _label(rows[measure_row][j + 1])) != ("#", "%"):
                raise ValueError(f"Table 14: unexpected measures under {c!r}")
            cols[TABLE14_BEDROOMS[_label(c)]] = (j, j + 1)
    flag_col = next(j for j, c in enumerate(rows[measure_row]) if _label(c) == "Metro/RV")
    if sorted(cols) != sorted(TABLE14_BEDROOMS.values()):
        raise ValueError(f"Table 14: unexpected header {rows[head]}")
    quarter = report_quarter(path)
    records = []
    for i in range(measure_row + 1, len(rows)):
        name = _label(rows[i][0])
        if not name:
            continue
        flag = _label(rows[i][flag_col])
        if name == "Total":
            row_type, metro = "victoria", None
        elif flag in ("M", "R"):
            row_type, metro = "lga", flag == "M"
        else:
            raise ValueError(f"Table 14: row {i} ({name!r}) has Metro/RV flag {flag!r}")
        for dwelling, (jc, js) in cols.items():
            records.append({"dwelling": dwelling, "lga": "Victoria" if row_type == "victoria" else clean_name(name),
                            "row_type": row_type, "metro": metro, "quarter": quarter,
                            "affordable_lettings": clean_value(rows[i][jc], f"Table 14!r{i}c{jc}"),
                            "affordable_share": clean_value(rows[i][js], f"Table 14!r{i}c{js}")})
    df = pd.DataFrame.from_records(records)
    cols_out = ["dwelling", "lga", "row_type", "metro", "quarter", "affordable_lettings", "affordable_share"]
    return _check_unique(df[cols_out], ["dwelling", "lga"])


def duplicate_mismatches(df: pd.DataFrame, keep: str, dup: str, measures: list[str]) -> pd.DataFrame:
    """Quarters where a repeated total row disagrees with the first one."""
    a = df[df["row_type"] == keep].set_index(["dwelling", "quarter"])[measures]
    b = df[df["row_type"] == dup].set_index(["dwelling", "quarter"])[measures]
    if not a.index.sort_values().equals(b.index.sort_values()):
        raise ValueError(f"'{dup}' rows cover different quarters from '{keep}' rows")
    joined = a.join(b.add_suffix("_repeat"))
    differs = pd.Series(False, index=joined.index)
    for m in measures:
        differs |= ~(joined[m].eq(joined[f"{m}_repeat"]) | (joined[m].isna() & joined[f"{m}_repeat"].isna()))
    return joined[differs].reset_index()


def _drop_duplicate(df: pd.DataFrame, keep: str, dup: str, tolerances: dict[str, float]) -> pd.DataFrame:
    """Drop repeated total rows after checking they match the first ones.

    The repeats are built from Metro + Non-Metro, so they can be off by a
    letting or two (or 0.001 of a share) from the first total, which equals
    the sum of the areas above it. Anything beyond `tolerances` raises.
    """
    bad = duplicate_mismatches(df, keep, dup, list(tolerances))
    for m, limit in tolerances.items():
        gap = (bad[m] - bad[f"{m}_repeat"]).abs()
        if (gap > limit + 1e-9).any() or (bad[m].isna() != bad[f"{m}_repeat"].isna()).any():
            raise ValueError(f"'{dup}' rows differ from '{keep}' rows in {m}:\n{bad.head()}")
    return df[df["row_type"] != dup].copy()


def _check_unique(df: pd.DataFrame, key: list[str]) -> pd.DataFrame:
    """Raise on duplicate keys; store counts as nullable integers."""
    dupes = df.duplicated(key, keep=False)
    if dupes.any():
        raise ValueError(f"duplicate keys:\n{df.loc[dupes, key].head()}")
    df = df.reset_index(drop=True)
    for col in ("lettings", "affordable_lettings"):
        if col in df:
            if not (df[col].dropna() % 1 == 0).all():
                raise ValueError(f"{col} has non-whole values")
            df[col] = df[col].astype("Int64")
    return df
